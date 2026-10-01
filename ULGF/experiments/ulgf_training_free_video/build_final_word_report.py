from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'runs' / 'final_word_report_20260930'
OUT.mkdir(exist_ok=True)
HANDOFF = ROOT / 'runs/final_diagnostic_handoff_20260930T215538548904Z'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.75)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(27)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = False
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = foot.add_run('ULGF experimental branch  |  ')
r.font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)

def clean(s):
    return s.replace('**','').replace('`','').replace('—','-')

def para(s, style=None):
    return doc.add_paragraph(clean(s), style)

def fig(path, caption, width=7):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set('descr', caption)
    para(caption, 'Caption')

def table(lines):
    rows = [[clean(v.strip()) for v in line.strip('|').split('|')] for line in lines if not re.match(r'^\|\s*---',line)]
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [1.25,2.65,3.1]
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,(cell,val) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]); cell.text=val
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'85'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DDEBF7' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF')); pr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                for r in p.runs: r.font.size=Pt(10); r.bold=i==0
        tr=t.rows[-1]._tr.get_or_add_trPr()
        tr.append(OxmlElement('w:cantSplit'))
        if i==0: tr.append(OxmlElement('w:tblHeader'))

def body(text):
    blocks=re.split(r'\n\s*\n',text.strip())
    for block in blocks:
        if block.startswith('|'): table(block.splitlines()); continue
        if re.match(r'^(\d+\. |\- )',block):
            items=re.split(r'\n(?=\d+\. |\- )',block)
            for item in items:
                s=' '.join(x.strip() for x in item.splitlines())
                s=re.sub(r'^\d+\. ','',s)
                s=re.sub(r'^- \[x\] ','Complete - ',s)
                s=re.sub(r'^- \[ \] ','Pending - ',s)
                s=re.sub(r'^- ','',s)
                para(s,'List Bullet')
        else: para(' '.join(block.splitlines()))

sections = re.split(r'^## ', (ROOT/'BRANCH_FINAL_REPORT.md').read_text(encoding='utf-8'),flags=re.M)[1:]
parts={s.split('\n',1)[0]:s.split('\n',1)[1] for s in sections}
doc.add_paragraph('ULGF training free video experiment', 'Title')
doc.add_paragraph('Final diagnostic report', 'Subtitle')
para('30 September 2026')
doc.add_heading('Decision and scope',1)
body(parts['Decision and scope'])
fig(HANDOFF/'sequence_5frame/contact_sheet.png','Figure 1. Accepted five-frame composition. The same selected fish pixels move over one supplied background plate.')
para('The deliverable links frame images, per-frame bounding boxes and persistent assigned identity. Its verified contribution is a reproducible compositing control, with explicit provenance and limitations.')

doc.add_page_break()
doc.add_heading('Evidence and progression',1)
body(parts['Evidence and progression'])
fig(ROOT/'reviews/fish_crop_20260929T015815740178Z/comparison.png','Figure 2. Fish-centred source crop, VAE reconstruction and image-conditioned candidates. Strength 0.05 was selected for the mask review.')

doc.add_page_break()
doc.add_heading('Mask and background review',1)
para('The approved silhouette follows the user-provided boundary, including the narrow tail and fins. It contains 4,932 foreground pixels. Earlier outlines included excess background below the gills and omitted parts of the tail.')
fig(ROOT/'reviews/fish_mask_user_v3/user_outline_comparison.png','Figure 3. Recorded mask review comparing the user outline with the derived boundary and earlier mask.')
para('The user confirmed that the cyan boundary followed the intended outline. A binary mask preserves these selected pixels but cannot separate translucent fins from background colour already mixed into the keyframe.')
fig(ROOT/'runs/reference_plate_cpu_20260929T223947096913Z/comparison.png','Figure 4. Earlier background repair above and the supplied clean plate below. Mask and motion remain fixed; the replacement plate is reused across the sequence.')
para('The clean plate is a supplied replacement, not recovered hidden scene geometry. The short composition and its playback were accepted for this diagnostic.')

doc.add_page_break()
doc.add_heading('Delivered sequences',1)
body(parts['Delivered sequences'])
fig(HANDOFF/'sequence_48frame/contact_sheet.png','Figure 5. Sampled positions from the 48-frame control. The fish is translated rigidly; no fin articulation or learned temporal motion is claimed.')

doc.add_page_break()
doc.add_heading('Validation results',1)
body(parts['Validation results'])
doc.add_heading('Identity provenance and evaluation boundaries',1)
body(parts['Identity, provenance and evaluation boundaries'])

doc.add_page_break()
doc.add_heading('Using the handoff',1)
body(parts['How to use the handoff'])
doc.add_heading('Closure checklist and next workstream',1)
body(parts['Closure checklist and next workstream'])
para('Evidence location: final_diagnostic_handoff_20260930T215538548904Z. The two sequence folders contain the validation records, source provenance and checksums supporting this report.')
doc.core_properties.title='ULGF training free video experiment final diagnostic report'
doc.core_properties.subject='Single object CPU compositing control and verified sequence deliverables'
doc.core_properties.author='ULGF Project'
dest=OUT/'ULGF_Final_Diagnostic_Report.docx'
doc.save(dest)
assert len(Document(dest).inline_shapes)==5
print('Created DOCX with five embedded figures and all report sections.')
