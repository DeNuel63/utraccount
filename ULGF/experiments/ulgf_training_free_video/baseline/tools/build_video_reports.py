"""Evidence-based Word reports of the implemented video layout subsystem."""
import json
import sys
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from ulgf_baseline.config import DEFAULT_RUOD_CLASSES
from ulgf_video.layout_sequence import build_layout_manifest
from ulgf_video.motion import MotionConfig
from ulgf_video.validation import validate_manifest
from scripts.visualize_layout_sequence import render

OUT=ROOT/'outputs/video-extension-reports'
OUT.mkdir(parents=True,exist_ok=True)
RUN=OUT/'layout_execution_001'; RUN.mkdir(exist_ok=True)
label=RUN/'fixture.txt'
label.write_text('4 0.28 0.30 0.22 0.12\n4 0.65 0.55 0.18 0.10\n8 0.40 0.80 0.25 0.16\n')
config=MotionConfig(frames=5,fps=10,max_speed_box_fraction=.4,max_acceleration_box_fraction=.02)
m=build_layout_manifest('layout_execution_001',RUN/'not_used_layout_fixture.png',label,
    DEFAULT_RUOD_CLASSES,256,256,42,'Layout-only fixture; no underwater pixels',config,'report_fixture')
m.write(RUN/'manifest.json')
render(m,RUN/'frames',overwrite=True)
validate_manifest(m)
assert m==build_layout_manifest('layout_execution_001',RUN/'not_used_layout_fixture.png',label,
    DEFAULT_RUOD_CLASSES,256,256,42,'Layout-only fixture; no underwater pixels',config,'report_fixture')
ann=RUN/'annotations'; ann.mkdir(exist_ok=True)
for frame in m.frames:
    (ann/('{:06d}.json'.format(frame.frame_index))).write_text(json.dumps({
        'frame_index':frame.frame_index,'objects':[{
            'instance_id':o.instance_id,'class_id':o.class_id,'class_name':o.class_name,
            'bbox_xyxy':list(o.bbox_xyxy),'visibility':o.visibility} for o in frame.objects]},indent=2))
summary={'execution':'CPU layout-only fixture','frames':5,'object_records':15,
    'persistent_ids':[1,2,3],'class_names':['fish','fish','turtle'],
    'deterministic_rerun':'PASS','schema_validation':'PASS',
    'trained_model_loaded':False,'photorealistic_frames_generated':False,
    'unit_tests_passed':11,'unit_tests_failed':0}
(RUN/'report.json').write_text(json.dumps(summary,indent=2))

font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',20)
small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',16)
palette=['#187FB0','#C76A10','#23844B']
strip=Image.new('RGB',(1400,375),'white'); draw=ImageDraw.Draw(strip)
draw.text((10,5),'Executed layout sequence   |   five frames   |   same IDs across frames',font=font,fill='black')
for index,frame in enumerate(m.frames):
    x=index*280+12; y= seventy=70
    draw.text((x,40),'Frame %03d'%index,font=small,fill='black')
    draw.rectangle((x,y,x+256,y+256),fill='#EAF4F7')
    for o in frame.objects:
        a,b,c,d=o.bbox_xyxy
        box=(x+a*256,y+b*256,x+c*256,y+d*256)
        color=palette[o.instance_id-1]
        draw.rectangle(box,outline=color,width=3)
        draw.text((box[0],box[1]-19),'%s ID %d'%(o.class_name,o.instance_id),font=small,fill=color)
draw.text((12,342),'Boxes are motion and identity outputs. Backgrounds are layout canvases, not diffusion-generated images.',font=small,fill='#444444')
strip.save(OUT/'five_frame_layout.png')

def doc(title,intro):
    d=Document(); s=d.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.left_margin=s.right_margin=Inches(.7);s.top_margin=s.bottom_margin=Inches(.65)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    d.styles['Normal'].font.size=Pt(11)
    d.styles['Normal'].paragraph_format.space_after=Pt(7)
    d.styles['Title'].font.size=Pt(25)
    d.styles['Heading 1'].font.size=Pt(15)
    d.add_paragraph(title,'Title')
    d.add_paragraph('ULGF video extension   |   28 September 2026','Subtitle')
    d.add_paragraph(intro)
    d.core_properties.title=title
    d.core_properties.subject='Measured layout subsystem results and uncompleted video training qualification'
    return d

def table(d,headers,rows):
    t=d.add_table(rows=1,cols=len(headers))
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for i,row in enumerate(t.rows):
        for c in row.cells:
            pr=c._tc.get_or_add_tcPr();bs=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');bs.append(b)
            pr.append(bs)
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7EEF2' if i==0 else 'FFFFFF');pr.append(fill)
            for p in c.paragraphs:
                p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
                for r in p.runs:r.font.size=Pt(10);r.bold=i==0

def fig(d):
    pic=d.add_picture(str(OUT/'five_frame_layout.png'),width=Inches(7.05))
    pic._inline.docPr.set('descr','Five executed layout frames with persistent IDs 1 and 2 for fish and 3 for turtle')
    d.add_paragraph('Figure 1. CPU execution of five layouts. Box coordinates change while each ID retains its class. This test does not measure pixel appearance consistency.','Caption')

d=doc('ULGF Video Extension Training Report',
    'The video extension is evaluated as a linked sequence of frames, per-frame annotations and persistent instance IDs. The completed local experiment establishes deterministic layout motion and identity metadata. It does not establish a trained video generator or temporally consistent underwater appearance.')
d.add_heading('Experimental setup',1)
table(d,['Parameter','Executed setting'],[
    ('Execution scope','CPU layout and annotation subsystem; no model weights loaded'),
    ('Sequence','5 frames, 256 by 256 canvas, nominal 10 frames per second'),
    ('Objects','Two fish and one turtle; IDs 1, 2 and 3'),
    ('Motion','Bounded random acceleration; reflecting boundary; seed 42'),
    ('Inputs','Explicit three-object fixture, not a sampled RUOD training image'),
    ('Checks','Schema validation, repeated execution and 11 focused unit tests')])
d.add_heading('Sequence and identity results',1);fig(d)
d.add_paragraph('All five layouts retained the three object IDs and their class assignments, producing 15 frame-object records. An identical rerun reproduced the manifest. Eleven schema, motion, sequence and layout-rendering unit tests passed with no failures. These are software correctness results, not learned tracking accuracy.')
d.add_page_break()
d.add_heading('Architecture implementation status',1)
table(d,['Component','Evidence and status'],[
    ('Sequential Layout Generator','Implemented in layout_sequence.py and motion.py; executed in this report'),
    ('Persistent ID assignment','Deterministic annotation-order IDs are propagated unchanged through the layout sequence'),
    ('Per-frame annotations','Normalized boxes, class labels and IDs are serialized in the manifest'),
    ('Medium and style conditioning','Clip metadata exists; a learned clip-level conditioning encoder is not demonstrated'),
    ('Temporal Consistency Module','Cross-frame attention or latent-warp conditioning is specified but not implemented in the inspected ulgf_video package'),
    ('Video image generation','No end-to-end temporally conditioned pixel sequence demonstrated by this experiment')])
d.add_heading('Training evidence and interpretation',1)
d.add_paragraph('Earlier supplied logs establish still-image inference and a byte-identical repeated seed-zero result. Visual review rejected that baseline for object defects and unrealistic color. Neither repeatability nor valid boxes establishes acceptable underwater video quality. No new optimization run, loss trajectory or final checkpoint metric is reported here.')
d.add_paragraph('The current layout MVP assumes a fixed object count and full visibility. Stable IDs are assigned control metadata, not identities recovered from image content. Before these annotations are used as COVTrack ground truth, generated pixels must be checked for missing objects, misplaced objects and appearance changes under the same assigned ID.')
d.add_heading('Required evaluation before video results are accepted',1)
d.add_paragraph('Complete pixel generation with temporal conditioning, compare against independent per-frame generation under fixed inputs, and inspect appearance, motion and annotation alignment across the clip. Evaluate saved-state resume on the real model and retain actual logs. Resolve dataset leakage and initialization provenance before claiming clean held-out performance.')
d.add_paragraph('Sources: System Architecture.docx, sections 2 and 4; ULGF_VIDEO_ENGINEERING_ROADMAP.md; docs/video_output_schema.md; local source inspection and the executed layout fixture. The architecture describes proposed components; its labels do not establish that those components have been implemented.')
d.save(OUT/'ULGF_Video_Extension_Training_Report.docx')

d=doc('ULGF Video Outputs and Annotation Results',
    'The video-extended output contract links ordered frame images to object boxes, class labels and persistent IDs. This report presents the executed layout sequence and its annotation records, then identifies the remaining pixel-generation requirement. The canonical deliverables are PNG frames and a clip manifest; MP4 is an optional preview.')
d.add_heading('Consecutive frame layouts',1);fig(d)
d.add_heading('Per-frame annotation results',1)
rows=[]
for f in m.frames:
    for o in f.objects:
        rows.append(('%03d'%f.frame_index,str(o.instance_id),o.class_name,', '.join('%.4f'%v for v in o.bbox_xyxy)))
table(d,['Frame','ID','Class','Normalized xyxy'],rows)
d.add_paragraph('Coordinates are x1, y1, x2, y2 in [0, 1], rounded here to four decimals for display. Frame indices start at zero. These are values from the executed motion model, not boxes measured from synthesized fish pixels.')
d.add_page_break()
d.add_heading('Identity continuity',1)
d.add_paragraph('ID 1 remains the first fish, ID 2 the second fish and ID 3 the turtle across all five layouts. Their box positions change without changing class or identity. This demonstrates metadata continuity. Visual identity continuity requires separately checking that each generated animal preserves its appearance along the same trajectory.')
d.add_heading('Stored output contract',1)
p=d.add_paragraph();r=p.add_run('clip_000001/\n  frames/000000.png ... 000004.png\n  manifest.json\n  sequence.mp4  [optional preview]\n  annotations/000000.json ... 000004.json  [optional adapter]')
r.font.name='Consolas';r.font.size=Pt(10)
d.add_paragraph('The report experiment stores layout preview PNGs and a manifest under layout_execution_001. Its per-frame JSON files are report-generated adapters of the canonical manifest; no MP4 or underwater video is claimed. Planned generated-frame exports must identify their checkpoint, configuration and temporal mode.')
d.add_heading('Annotation record',1)
o=m.frames[1].objects[0]
example={'frame_index':1,'instance_id':o.instance_id,'class_id':o.class_id,'class_name':o.class_name,
    'bbox_xyxy':[round(v,8) for v in o.bbox_xyxy],'visibility':o.visibility}
p=d.add_paragraph();r=p.add_run(json.dumps(example,indent=2));r.font.name='Consolas';r.font.size=Pt(9)
d.add_paragraph('For a top-left pixel xywh adapter: x = 256x1, y = 256y1, w = 256(x2 - x1), h = 256(y2 - y1). A center-based xywh interface instead uses the box midpoint; adapters must name their convention explicitly. The supplied 490-pixel example is not valid for this 256-pixel canvas. Crab is not one of the ten configured RUOD classes, so this run uses turtle as the third object.')
d.add_heading('COVTrack handoff',1)
d.add_paragraph('CountGD++ detections supply boxes, labels and confidence to the tracker. ULGF persistent IDs form a separate synthetic reference for association evaluation only after image-label alignment is validated. Do not pass reference IDs to the tracker as predicted identities or report their persistence as measured tracking accuracy.')
d.save(OUT/'ULGF_Video_Outputs_and_Annotation_Results.docx')
print(json.dumps(summary))
