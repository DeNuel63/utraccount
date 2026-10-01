"""Create projected visual briefs separately from measured training outputs."""
from pathlib import Path
import math
import shutil
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/projected-results-documents'
OUT.mkdir(parents=True,exist_ok=True)
SOURCE=Path('C:/Users/De_Nuel 🔥/.codex/generated_images/01a0c0b9-58ea-7a41-b1d8-2c2e37f54b74/exec-8b3500b6-bd9c-46e8-9c88-783d0a6e76a2.png')
shutil.copy2(SOURCE,OUT/'underwater_target_concepts.png')

def document(title,opening):
    doc=Document()
    s=doc.sections[0]
    s.page_width=Inches(8.5); s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.65)
    s.left_margin=s.right_margin=Inches(.75)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(26)
    doc.styles['Heading 1'].font.size=Pt(16)
    doc.styles['Heading 2'].font.size=Pt(12)
    doc.add_paragraph(title,'Title')
    doc.add_paragraph('ULGF research project   |   28 September 2026','Subtitle')
    doc.add_paragraph(opening)
    doc.core_properties.title=title
    doc.core_properties.subject='Projected outcomes and quality targets, not completed training results'
    doc.core_properties.author='ULGF project'
    return doc

def picture(doc,path,width,caption):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
    run=p.add_run(); image=run.add_picture(str(path),width=Inches(width))
    image._inline.docPr.set('descr',caption)
    doc.add_paragraph(caption,'Caption')

def table(doc,head,rows):
    t=doc.add_table(rows=1,cols=len(head)); t.autofit=False
    for c,h in zip(t.rows[0].cells,head): c.text=h
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)
    for ri,row in enumerate(t.rows):
        for c in row.cells:
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                b=OxmlElement('w:'+edge); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
            pr.append(borders)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EEF2' if ri==0 else 'FFFFFF'); pr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(5)
                for r in p.runs: r.font.size=Pt(10); r.bold=ri==0
    return t

# A schematic has no numeric loss scale or claimed epoch count.
im=Image.new('RGB',(1400,680),'white'); draw=ImageDraw.Draw(im)
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',25)
small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',22)
draw.text((95,25),'Projected learning pattern',font=font,fill='black')
draw.text((95,65),'Qualitative trajectories only — not recorded losses',font=small,fill='#555555')
draw.line([(110,135),(110,550),(1310,550)],fill='#303030',width=3)
draw.text((115,108),'Higher loss',font=small,fill='#555555')
draw.text((115,558),'Early optimization',font=small,fill='#555555')
draw.text((1070,558),'Later optimization',font=small,fill='#555555')
for color,offset,floor in [('#127484',0,420),('#B6742A',20,355)]:
    points=[]
    for i in range(1100):
        x=150+i; y=floor-(floor-180-offset)*math.exp(-i/255)+5*math.sin(i/29)
        points.append((x,int(y)))
    draw.line(points,fill=color,width=5)
draw.line([(720,105),(780,105)],fill='#127484',width=5)
draw.text((790,89),'Training',font=small,fill='black')
draw.line([(1000,105),(1060,105)],fill='#B6742A',width=5)
draw.text((1070,89),'Validation',font=small,fill='black')
draw.text((110,620),'A widening gap may indicate overfitting; lower loss alone does not prove better images.',font=small,fill='#555555')
im.save(OUT/'projected_learning_pattern.png')

doc=document('ULGF Projected Underwater Outputs',
    'Our visual target is recognizable underwater objects with plausible anatomy, restrained color and coherent surroundings. The concepts below were generated independently for this preview; they are not ULGF checkpoint outputs or a guarantee of achievable quality.')
picture(doc,OUT/'underwater_target_concepts.png',5.6,
    'Visual targets clockwise from upper left: fish, cuttlefish, turtle, and starfish with echinus. Independent AI generated concepts, not measured model results.')
doc.add_heading('What successful ULGF output should demonstrate',1)
doc.add_paragraph('The requested class should be recognizable, occupy its intended layout region and remain distinct from the background. Fish should retain coherent fins and tails; cuttlefish should have plausible mantle and tentacle structure. Reef texture and water color should not produce the severe distortions seen in the earlier baseline.')
doc.add_paragraph('These high resolution concepts define a visual direction. The planned 256 by 256 ULGF outputs will contain less detail; actual anatomy, layout fidelity and scene diversity must be assessed after training.')
doc.save(OUT/'ULGF_Projected_Underwater_Outputs.docx')

doc=document('ULGF Projected Training Results',
    'The desired outcome is a reproducible checkpoint that improves class appearance and scene realism while respecting input layouts. This report presents the expected form of the evaluation, not a completed training result. No final accuracy, loss, FID or training duration has been measured.')
doc.add_heading('Verified starting point',1)
table(doc,['Item','Latest reported evidence'],[
    ('Reviewed v11 membership','7,515 training images; 842 validation images; 4,200 official test images'),
    ('Baseline inference','A 256 by 256 RGB image was generated; duplicate seed zero runs had identical output hashes'),
    ('Baseline quality','User review rejected object appearance, color and scene artifacts'),
    ('Training status','Full training and real model resume qualification remain uncompleted')])
doc.add_heading('Projected optimization behavior',1)
picture(doc,OUT/'projected_learning_pattern.png',6.7,
    'Expected curve shape under successful optimization. Axes intentionally omit invented numerical loss values and epoch counts.')
doc.add_paragraph('Training and held out validation losses may decrease and then flatten. Checkpoint selection should combine fixed validation measurements with visual inspection; continued training loss improvement is insufficient if validation quality deteriorates.')
doc.add_page_break()
doc.add_heading('Projected visual improvement',1)
picture(doc,OUT/'underwater_target_concepts.png',4.0,
    'Independent visual targets for the evaluation panel. These are not before and after training samples.')
table(doc,['Evaluation dimension','Desired observation'],[
    ('Class appearance','Recognizable requested animals without severe structural defects'),
    ('Layout adherence','Requested objects lie within or near their specified regions'),
    ('Color and context','Plausible underwater attenuation and coherent reef or seabed detail'),
    ('Diversity','Different seeds change scenes without repeatedly copying training examples'),
    ('Resume integrity','A saved training state reproduces the next controlled update within a defined tolerance')])
doc.add_heading('Evidence required for the final results report',1)
doc.add_paragraph('Replace the projected figures with fixed prompt and seed comparisons, actual training and validation logs, the selected checkpoint hash, GPU memory measurements and a verified resume comparison. Record failures as well as successful examples. Resolve the outstanding leakage review and initialization provenance before claiming an uncontaminated held out evaluation.')
doc.add_paragraph('Basis: project supplied v11 application report and baseline inference logs. Split membership remains provisional until leakage clearance; none of the visual concepts establishes a model performance score.')
doc.save(OUT/'ULGF_Projected_Training_Results.docx')
print('Created two Word documents in outputs/projected-results-documents')
