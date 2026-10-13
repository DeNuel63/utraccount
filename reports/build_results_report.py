from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER

OUT = Path(__file__).resolve().parent
NAME = 'Utraccount Training and Execution Results'
pages = [
[
('title', NAME),
('sub', 'Project results report   |   12 October 2026'),
('h', 'Overall result'),
('p', 'The Utraccount pipeline connected ULGF generation, CountGD++ detection and COVTrack tracking on a five-frame sequence. The user executed the models on a Google Colab Tesla T4 and supplied logs and saved-output summaries. Execution completed through retries that reused generated frames and detections.'),
('p', 'The project also records a separate ULGF training continuation of 100 optimizer updates on 7 October 2026. Its expanded paired evaluation did not demonstrate improvement. No new CountGD++ or COVTrack training run is established by the evidence reviewed for this report.'),
('h', 'Status by model'),
('table', (['Model', 'Training evidence', 'Execution evidence'], [
['ULGF', '100-update continuation recorded', 'Five frames generated'],
['CountGD++', 'No new training verified', 'Detections exported for five frames'],
['COVTrack', 'No new training verified', 'Tracking outputs for five frames'],
])),
('h', 'What this result establishes'),
('p', 'The sequence passed through all three stages in order. COVTrack consumed the external CountGD++ detection records and produced tracking records. Predicted ID 0 appears in every frame; predicted ID 1 appears on the additional detection in frame 3.'),
('p', 'This is a functional connection result. Detection accuracy, correct physical identity, temporal realism and benchmark tracking performance were not measured. There was no joint training of the three models.'),
('h', 'Evidence basis'),
('p', 'Training figures are taken from the repository training report. Runtime versions, errors and final frame counts are taken from user-supplied Colab output in this conversation. The final Colab JSON files and images were not downloaded for an independent audit. This report consolidates the available evidence rather than certifying a complete experimental benchmark.'),
],
[
('h', 'ULGF training results'),
('p', 'Recorded run 20261007T103836275651Z continued existing RUOD-trained still-image weights for 100 updates. Training, inference export, backup and an actual full-state resume were exercised. The run was an engineering continuation, not a full training campaign.'),
('table', (['Setting', 'Recorded value'], [
['Resolution', '256 by 256 pixels'],
['Optimizer and learning rate', 'AdamW8bit; 5e-6; constant schedule'],
['Updates and batch size', '100 updates; batch size 1; accumulation 1'],
['Trainable component', 'UNet; VAE and text encoder frozen'],
['Precision and seed', 'fp16; training seed 42'],
['Engineering subset', '7,215 training and 812 validation images'],
['Training loss', 'Weighted epsilon MSE; foreground weight 1.0'],
])),
('h', 'Paired evaluation'),
('p', 'Lower denoising loss is better. Original and step-100 checkpoints were compared using matched inputs and stochastic settings, with no optimizer updates during evaluation.'),
('table', (['Evaluation', 'Original', 'Step 100', 'Change'], [
['4 fixed examples', '0.12232011', '0.12222775', '0.0755% lower'],
['64 images and 3 seeds', '0.16379158', '0.16394605', '0.0943% higher'],
])),
('p', 'The expanded evaluation comprised 192 paired probes. Only 15 of 64 images improved after averaging seeds, and all three seed-level mean losses increased. The mean after-minus-before difference was 0.00015448; its conditional image-bootstrap 95% interval was 0.00009478 to 0.00021819.'),
('p', 'Meaningful quality improvement was not demonstrated. Denoising loss is not a perceptual metric. Holdout provenance and remaining leakage review prevent a clean independent-holdout claim.'),
('h', 'Checkpoint and resume evidence'),
('p', 'The repository records verified inference export and backup plus a successful full-state resume from step 100 to 101. Peak allocated GPU memory in the resume test was 10,491.25 MiB. Equivalence to uninterrupted training was not tested; inference weights alone do not preserve optimizer state. The checkpoint used for the October 12 pipeline run has not been linked by hash to this step-100 export.'),
],
[
('h', 'Generation and model connections'),
('p', 'ULGF generated five frames in the October 12 execution. Its completion message was INFERENCE_COMPLETE_REVIEW_REQUIRED. The requested generation configuration used seed 0, 100 steps and guidance 5.0; the final resolved configuration should be retained to confirm these settings for reproduction.'),
('h', 'ULGF to CountGD++'),
('p', 'The handoff uses ordered RGB images, a sequence identifier and zero-based frame indices. CountGD++ processes the generated images using the selected text class. The prepared configuration uses fish and a confidence threshold of 0.23. The final detection records remain the authority for actual run settings.'),
('h', 'CountGD++ to COVTrack'),
('p', 'CountGD++ exports confidence scores and original-image pixel boxes in centre x, centre y, width, height format. The tracking stage converts them into left, top, right, bottom, score records and shared integer class IDs. COVTrack reuses its RoI feature heads and association logic for these external boxes.'),
('p', 'The integration keeps requested ULGF reference annotations separate from predicted detections and tracking IDs. Requested generation identities are not used as tracking ground truth. The custom vocabulary includes ten RUOD classes with fish at class ID 4.'),
('h', 'Earlier generation diagnostic'),
('p', 'A separate September 30 handoff documented five-frame and 48-frame translation compositions at 256 by 256 pixels. Those experiments used a source-derived fish keyframe, a reviewed mask and a supplied background. Their identities were assigned by the composition pipeline. They are distinct from the October 12 model connection test and do not constitute learned video generation or tracker evaluation.'),
('h', 'Training and performance boundaries'),
('p', 'CountGD++ and COVTrack were used for inference with existing checkpoints. Drive training-result folder names alone do not establish training completion or scores. No verified CountGD++ training loss, counting MAE or RMSE, detection precision or recall, or COVTrack TETA, HOTA, MOTA or IDF1 result is available in the reviewed evidence.'),
('p', 'The earlier ULGF report retained visual defects and did not validate temporal consistency. Stable exported tracking IDs in this smoke sequence do not resolve those generation-quality questions.'),
],
[
('h', 'Five frame execution results'),
('p', 'The user read detection_sequence.json and tracking_sequence.json from the retry output. Their check required the same sequence identifier, matching frame indices and five frames in each file. The printed result is reproduced below.'),
('table', (['Frame index', 'Detections', 'Tracks', 'Predicted IDs'], [
['0', '1', '1', '0'], ['1', '1', '1', '0'], ['2', '1', '1', '0'],
['3', '2', '2', '0 and 1'], ['4', '1', '1', '0'],
])),
('p', 'Across the five frames, the output contains six detection instances and six tracking records, with two distinct predicted IDs. ID 0 is present in all five frames. ID 1 is present only in frame 3. These counts do not establish whether the extra detection represents a real object, a duplicate or a false positive.'),
('h', 'Execution recovery'),
('p', 'The first attempt completed ULGF generation, then stopped when CountGD++ inherited Colab\'s inline Matplotlib backend. Setting MPLBACKEND to Agg allowed the detector to proceed. COVTrack then required lap, BDD100K, TETA and filelock because of imports in its package initialization.'),
('p', 'COVTrack next reached checkpoint loading but its image preprocessing raised KeyError for frame_id. The inference entry point was corrected to supply the frame index at the top level as well as inside img_info. The subsequent tracking attempt produced the five-frame output summarized above.'),
('h', 'Completion status'),
('p', 'All three stages completed on the same generated sequence across resumed attempts. A single uninterrupted successful launcher invocation was not recorded. Tracking was finally invoked directly, so the earlier launcher status.json still records its historical failure; it has not been rewritten as successful.'),
('h', 'Result location'),
('p', 'Google Drive output folder: My Drive/connected_pipeline_results/retry_20261012_224051. Key evidence files are resolved_config.json, detection_sequence.json, per-frame detection JSON, tracking_sequence.json and per-frame tracking JSON. The original generation attempt used run_20261012_223233. Generated images remain under the configured ULGF experiment run folder in the Colab project copy unless separately backed up.'),
],
[
('h', 'Verified runtime and final checklist'),
('p', 'The run used a Tesla T4 with approximately 15 GB GPU memory. Colab\'s notebook kernel was Python 3.13.15; the model stages used separate interpreters. The system driver reported CUDA 13.0, while the model PyTorch builds used older CUDA runtimes.'),
('table', (['Stage', 'Python', 'PyTorch', 'Key packages'], [
['ULGF', '3.7.16', '1.12.1+cu113', 'Diffusers 0.4.1; Transformers 4.24.0; Accelerate 0.20.3'],
['CountGD++', '3.10.22', '2.1.2+cu118', 'torchvision 0.16.2; Transformers 4.40.2; Detectron2 0.6'],
['COVTrack', '3.8.20', '1.10.0+cu113', 'MMCV 1.4.4; MMDetection 2.23.0; CLIP'],
])),
('p', 'CountGD++ and COVTrack import probes passed with CUDA available. CountGD++ GPU arithmetic passed, and COVTrack compiled GPU NMS returned the expected retained box. ULGF completed real generation. The supplied notebook transcript also records all six CountGD++ attention extension tests passing, including forward comparisons and numerical gradient checks.'),
('h', 'Completed connection checklist'),
('p', '[x] Confirm compatible environments; use separate environments if their dependencies conflict.'),
('p', '[x] Execute the three stages in order.'),
('p', '[x] Run one short sequence through all three stages.'),
('p', '[x] Confirm COVTrack receives CountGD++ detections and produces tracking outputs.'),
('h', 'Evidence references'),
('p', '1. ULGF/docs/TRAINING_RESULTS.md: October 7 training configuration, evaluation, backup and resume record.'),
('p', '2. ULGF/experiments/ulgf_training_free_video/BRANCH_FINAL_REPORT.md: earlier composition diagnostic and limitations.'),
('p', '3. integration/README.md, RUNTIME_SETUP.md and stage implementations: execution contracts and external-detection connection.'),
('p', '4. User-supplied October 12 Colab logs and final five-frame detection and tracking summary: actual execution evidence. Earlier CountGD++ component status documents predate this connection and are not current execution evidence.'),
('h', 'Overall conclusion'),
('p', 'Utraccount has a verified functional connection from generation to detection to tracking on one short sequence. Its recorded ULGF training continuation demonstrated engineering execution and recovery, but no quality gain. Broader accuracy, temporal quality and production readiness remain unmeasured.'),
],
]

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.75)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.06
for name in ['Title', 'Heading 1']:
    doc.styles[name].font.color.rgb = RGBColor(0,0,0)
doc.styles['Title'].font.size = Pt(26)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Utraccount results  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyReport', fontName='Helvetica', fontSize=10.5, leading=14, spaceAfter=7))
styles.add(ParagraphStyle(name='HeadingReport', fontName='Helvetica-Bold', fontSize=15, leading=18, spaceBefore=10, spaceAfter=7))
styles.add(ParagraphStyle(name='TitleReport', fontName='Helvetica-Bold', fontSize=25, leading=29, spaceAfter=12))
styles.add(ParagraphStyle(name='SubReport', fontName='Helvetica', fontSize=10, leading=14, textColor=colors.HexColor('#526174'), spaceAfter=12))
styles.add(ParagraphStyle(name='CellReport', fontName='Helvetica', fontSize=9, leading=12))
story = []
for page_no, blocks in enumerate(pages):
    if page_no:
        doc.add_page_break(); story.append(PageBreak())
    for kind, value in blocks:
        if kind == 'table':
            headers, rows = value
            n = len(headers)
            widths = ([1.25, 2.1, 3.65] if n==3 and headers[0]=='Model' else
                      [1.7, 5.3] if n==2 else
                      [2.3, 1.3, 1.3, 2.1] if headers[0]=='Evaluation' else
                      [1.3, 1.3, 1.3, 3.1] if headers[0]=='Frame index' else [1.0, .8, 1.4, 3.8])
            table = doc.add_table(rows=1, cols=n)
            table.autofit=False
            for cell, text, width in zip(table.rows[0].cells, headers, widths):
                cell.text=text; cell.width=Inches(width)
            for row in rows:
                cells=table.add_row().cells
                for cell, text, width in zip(cells,row,widths):
                    cell.text=text; cell.width=Inches(width)
            for idx,row in enumerate(table.rows):
                trpr=row._tr.get_or_add_trPr()
                if idx==0: trpr.append(OxmlElement('w:tblHeader'))
                for cell in row.cells:
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr=cell._tc.get_or_add_tcPr()
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'203B56' if idx==0 else ('F1F5F8' if idx%2 else 'FFFFFF')); tcpr.append(shade)
                    borders=OxmlElement('w:tcBorders')
                    for side in ['top','left','bottom','right']:
                        edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'D9D9D9'); borders.append(edge)
                    tcpr.append(borders)
                    margins=OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        m=OxmlElement('w:'+side);m.set(qn('w:w'),'90');m.set(qn('w:type'),'dxa');margins.append(m)
                    tcpr.append(margins)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(2)
                        for run in p.runs:
                            run.font.size=Pt(10)
                            if idx==0: run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
            data=[[Paragraph(escape(str(x)),styles['CellReport']) for x in r] for r in [headers]+rows]
            pdf_table=Table(data,colWidths=[w*72 for w in widths],repeatRows=1,hAlign='LEFT')
            pdf_table.setStyle(TableStyle([
                ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCE6EF')),
                ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F1F5F8')]),
                ('GRID',(0,0),(-1,-1),.4,colors.HexColor('#D9D9D9')),
                ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
                ('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),
            ]))
            story.extend([pdf_table,Spacer(1,9)])
        else:
            ds={'title':'Title','h':'Heading 1','sub':'Subtitle','p':'Normal'}[kind]
            doc.add_paragraph(value,style=ds)
            ps={'title':'TitleReport','h':'HeadingReport','sub':'SubReport','p':'BodyReport'}[kind]
            story.append(Paragraph(escape(value),styles[ps]))
doc.save(OUT/(NAME+'.docx'))

def footer_pdf(canvas, document):
    canvas.setFont('Helvetica',9); canvas.setFillColor(colors.HexColor('#526174'))
    canvas.drawString(54,30,'Utraccount Training and Execution Results')
    canvas.drawRightString(558,30,str(document.page))

pdf = SimpleDocTemplate(str(OUT/(NAME+'.pdf')),pagesize=(612,792),
    leftMargin=54,rightMargin=54,topMargin=48,bottomMargin=48,
    title=NAME,author='Utraccount project')
pdf.build(story,onFirstPage=footer_pdf,onLaterPages=footer_pdf)
print('Created Word and PDF reports')
